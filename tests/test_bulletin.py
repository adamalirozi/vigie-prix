import io

import pytest
from docx import Document

import config as cfg
from services import alerts as engine
from services import bulletin as B
from services import meb


@pytest.fixture(scope="module")
def data(replay_models, basket):
    ms = replay_models
    m = ms["Mil"]
    raw = engine.generate(ms, basket, "Référence", 3, 20, 15_000)
    raw = raw[(raw["Produit"] == "Mil") & (raw["Province"] == "Hadjer-Lamis")]
    validated = raw.assign(Valide_par="Analyste SISAAP", Date_decision="2026-10-22T10:00:00", Commentaire="")
    return B.BulletinData(
        commodity="Mil", scenario="Référence", horizon=3, last_obs=m.last_obs, dates_hist=m.dates_hist,
        dates_proj=m.dates_proj, table=m.province_table("Référence", 3), province="Hadjer-Lamis",
        series=m.province_series("Hadjer-Lamis", "Référence"), validated=validated,
        meb_table=meb.cost_table(ms, basket, "Référence", 3), meb_seuil=15_000, seuil_pct=20, replay=True)


def test_text_sections_are_complete_and_cautious(data):
    sec = B.sections(data)
    assert len(sec) == 5 and any(k.startswith("3. Alertes validées") for k in sec)
    text = B.build_text(data)
    assert "BROUILLON À VALIDER" in text and "mars 2024" in text and "Hadjer-Lamis" in text
    assert "non validé" in text and "précision n'est pas garantie" in text


def test_without_validated_alerts_the_section_says_so(data):
    empty = B.BulletinData(**{**data.__dict__, "validated": data.validated.iloc[0:0]})
    assert "Aucune alerte n'a été validée" in " ".join(B.sections(empty)["3. Alertes validées"])


def test_docx_opens_and_contains_banner_and_images(data):
    raw = B.build_docx(data)
    doc = Document(io.BytesIO(raw))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "BROUILLON À VALIDER" in text and "Mode relecture" in text
    assert len(doc.inline_shapes) == 2 and len(doc.tables) == 1  # deux graphiques, un tableau
    assert "BROUILLON" in doc.sections[0].footer.paragraphs[0].text


def test_charts_are_png(data):
    for buf in (B.chart_deviation(data), B.chart_trend(data)):
        assert buf.read(8) == b"\x89PNG\r\n\x1a\n"
    assert B.chart_trend(B.BulletinData(**{**data.__dict__, "series": None})) is None


def test_demo_banner_is_visible_in_config():
    assert "non validés par le SISAAP" in cfg.DEMO_BANNER
