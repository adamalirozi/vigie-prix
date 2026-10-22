"""Brouillon de rubrique « Projections de prix » pour le bulletin Al Aswak.

Le texte est généré à partir des tables de l'application ; il est un **brouillon à relire** : aucune phrase n'est
publiable sans validation d'un analyste du SISAAP. Export Word (avec graphiques) et texte brut.
"""

from __future__ import annotations

import io
from dataclasses import dataclass

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402

import config as cfg  # noqa: E402
from ui.formatting import kg, month_label, month_label_long, pct  # noqa: E402

BLUE, ORANGE, RED, GREY = "#4e79a7", "#f28e2b", "#e15759", "#aab2bd"


@dataclass
class BulletinData:
    commodity: str
    scenario: str
    horizon: int
    last_obs: pd.Timestamp
    dates_hist: list
    dates_proj: list
    table: pd.DataFrame  # province_table
    province: str
    series: dict | None  # province_series de la province sélectionnée
    validated: pd.DataFrame  # alertes validées
    meb_table: pd.DataFrame | None
    meb_seuil: float
    seuil_pct: float
    replay: bool = False


# --------------------------------------------------------------------------- #
# Texte
# --------------------------------------------------------------------------- #
def _names(df: pd.DataFrame, n: int) -> str:
    items = [f"{r.Province} ({kg(r.Prix_Actuel)})" for r in df.head(n).itertuples()]
    return ", ".join(items)


def sections(d: BulletinData) -> dict[str, list[str]]:
    t = d.table
    product = d.commodity.lower()
    end = month_label_long(d.dates_proj[d.horizon - 1])
    med_price, med_dev = t["Prix_Actuel"].median(), t["Ecart_Moy5_pct"].median()
    med_yoy, med_var = t["YoY_pct"].median(), t["Variation_pct"].median()
    watch = t[t["Variation_pct"] >= cfg.ALERT_MID * d.horizon / cfg.HORIZON].sort_values("Variation_pct", ascending=False)
    above = t[t["Ecart_Moy5_pct"] > d.seuil_pct]

    situation = [
        f"En {month_label_long(d.last_obs)}, le prix médian du {product} sur les {len(t)} provinces suivies est de "
        f"{kg(med_price)}. Il est inférieur de {pct(abs(med_dev), signed=False)} à la moyenne des cinq dernières "
        f"années pour le même mois" if med_dev < 0 else
        f"En {month_label_long(d.last_obs)}, le prix médian du {product} sur les {len(t)} provinces suivies est de "
        f"{kg(med_price)}, soit {pct(med_dev, signed=False)} au-dessus de la moyenne des cinq dernières années "
        "pour le même mois",
        f"et {'inférieur' if med_yoy < 0 else 'supérieur'} de {pct(abs(med_yoy), signed=False)} à son niveau d'il y a un an.",
        f"Les prix les plus élevés sont observés à : {_names(t.sort_values('Prix_Actuel', ascending=False), 3)}. "
        f"Les plus bas : {_names(t.sort_values('Prix_Actuel'), 3)}.",
    ]
    situation = [" ".join(situation[:2]), situation[2]]
    if len(above):
        situation.append(f"{len(above)} province(s) dépassent de plus de {d.seuil_pct:.0f} % leur moyenne quinquennale "
                         f"(seuil de démonstration, non validé) : {', '.join(above['Province'].head(6))}.")

    projection = [
        f"Selon la référence saisonnière (niveau 1, sans facteurs externes), la variation médiane projetée d'ici "
        f"{end} est de {pct(med_var)} (scénario « {d.scenario} »)."
    ]
    if len(watch):
        projection.append(
            f"{len(watch)} province(s) présentent une hausse projetée supérieure ou égale au repère de vigilance "
            f"modérée : {', '.join(f'{r.Province} ({pct(r.Variation_pct)})' for r in watch.head(6).itertuples())}."
        )
    else:
        projection.append("Aucune province ne présente de hausse projetée atteignant le repère de vigilance modérée.")
    projection.append("Ces projections sont indicatives : elles ne tiennent pas compte des chocs (récolte, conflit, "
                      "carburant, frontières) et leur précision n'est pas garantie (voir page « Fiabilité »).")

    alerts = []
    if d.validated.empty:
        alerts.append("Aucune alerte n'a été validée par un analyste à ce jour : cette rubrique reste vide tant que "
                      "le SISAAP n'a pas validé de projection.")
    else:
        for r in d.validated.itertuples():
            alerts.append(f"{r.Type} — {r.Produit}, {r.Province} : premier dépassement prévu en "
                          f"{month_label_long(r.Premier_Mois)} ; confiance {r.Confiance.lower()} ; "
                          f"validée par {r.Valide_par}.")

    meb = []
    if d.meb_table is not None and not d.meb_table.empty:
        m = d.meb_table
        over = m[m["Cout_Projete"] > d.meb_seuil]
        meb.append(
            f"Composante céréales du MEB (ration de {cfg.HOUSEHOLD_SIZE} personnes) : coût médian de "
            f"{m['Cout_Actuel'].median():,.0f} FCFA par mois, projeté à {m['Cout_Projete'].median():,.0f} FCFA d'ici "
            f"{end}.".replace(",", "\u202f")
        )
        meb.append(
            f"{len(over)} province(s) dépassent le seuil de {d.meb_seuil:,.0f} FCFA (seuil de démonstration, non validé)"
            .replace(",", "\u202f") + (f" : {', '.join(over['Province'].head(8))}." if len(over) else ".")
        )
    method = [
        "Source : WFP Price Database (HDX), prix de détail au kilo sur les marchés. Projection : dernier prix × rapport "
        "saisonnier × dérive du scénario ; intervalle à 80 % issu d'un rétro-test. Brouillon généré automatiquement ; "
        "à relire et valider avant toute diffusion.",
    ]
    return {"1. Situation actuelle": situation, f"2. Projections à {d.horizon} mois": projection,
            "3. Alertes validées": alerts, "4. MEB (composante céréales)": meb, "5. Méthode et limites": method}


def build_text(d: BulletinData) -> str:
    out = [f"AL ASWAK — BROUILLON À VALIDER — {d.commodity} — {month_label_long(d.last_obs)}", ""]
    for title, paragraphs in sections(d).items():
        out += [title, *paragraphs, ""]
    return "\n".join(out)


# --------------------------------------------------------------------------- #
# Graphiques (matplotlib : pas de navigateur requis)
# --------------------------------------------------------------------------- #
def _png(fig) -> io.BytesIO:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=160, bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf


def chart_deviation(d: BulletinData) -> io.BytesIO:
    t = d.table.dropna(subset=["Ecart_Moy5_Projete_pct"]).sort_values("Ecart_Moy5_Projete_pct")
    fig, ax = plt.subplots(figsize=(7, max(3.2, 0.26 * len(t) + 1)))
    colors = [RED if v > d.seuil_pct else BLUE for v in t["Ecart_Moy5_Projete_pct"]]
    ax.barh(t["Province"], t["Ecart_Moy5_Projete_pct"], color=colors)
    ax.axvline(0, color="#30343b", lw=1)
    ax.axvline(d.seuil_pct, color=RED, lw=1, ls="--")
    ax.set_xlabel(f"Écart du prix projeté ({month_label(d.dates_proj[d.horizon - 1])}) à la moyenne 5 ans (%)")
    ax.spines[["top", "right"]].set_visible(False)
    ax.tick_params(labelsize=8)
    return _png(fig)


def chart_trend(d: BulletinData) -> io.BytesIO | None:
    if d.series is None:
        return None
    s = d.series
    xh = list(range(len(d.dates_hist)))
    xp = [xh[-1] + i for i in range(0, cfg.HORIZON + 1)]
    last = next((v for v in reversed(s["hist"]) if pd.notna(v)), s["proj"][0])
    fig, ax = plt.subplots(figsize=(7, 3.4))
    lo = [last] + [b[0] for b in s["band"]]
    hi = [last] + [b[1] for b in s["band"]]
    ax.fill_between(xp, lo, hi, color=ORANGE, alpha=0.18, label="Intervalle à 80 %")
    ax.plot(range(len(s["ref"])), s["ref"], color=GREY, ls=":", lw=2, label="Moyenne 5 ans")
    ax.plot(xh, s["hist"], color=BLUE, lw=2.2, marker="o", ms=3, label="Prix observé")
    ax.plot(xp, [last, *s["proj"]], color=ORANGE, lw=2.2, ls="--", marker="o", ms=3, label="Projection")
    if s.get("actual"):
        ax.plot(xp[1:], s["actual"], color="#59a14f", lw=2.2, marker="D", ms=5, label="Réalisé")
    labels = [month_label(x) for x in d.dates_hist + d.dates_proj]
    ax.set_xticks(range(0, len(labels), 3), [labels[i] for i in range(0, len(labels), 3)], fontsize=8)
    ax.set_ylabel("FCFA / kg")
    ax.set_title(f"{d.commodity} · {d.province}", fontsize=10, loc="left")
    ax.legend(fontsize=7, frameon=False, ncol=3, loc="upper left")
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="y", color="#e4e7eb")
    return _png(fig)


# --------------------------------------------------------------------------- #
# Word
# --------------------------------------------------------------------------- #
def build_docx(d: BulletinData) -> bytes:
    doc = Document()
    for style in ("Normal",):
        doc.styles[style].font.name = "Calibri"
        doc.styles[style].font.size = Pt(10.5)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.9)
    sec.footer.paragraphs[0].text = (f"BROUILLON À VALIDER · Vigie Prix · {cfg.DEMO_BANNER}")
    sec.footer.paragraphs[0].runs[0].font.size = Pt(8)

    title = doc.add_heading("Al Aswak — Projections de prix", level=0)
    title.runs[0].font.size = Pt(22)
    sub = doc.add_paragraph(f"{d.commodity} · situation à {month_label_long(d.last_obs)} · horizon M+{d.horizon} · "
                            f"scénario « {d.scenario} »")
    sub.runs[0].font.color.rgb = RGBColor(0x5F, 0x67, 0x72)
    banner = doc.add_paragraph()
    run = banner.add_run("BROUILLON À VALIDER — non validé par le SISAAP. Données publiques WFP (démonstration). "
                         "Ne pas diffuser.")
    run.bold = True
    run.font.color.rgb = RGBColor(0xB4, 0x6A, 0x14)
    if d.replay:
        doc.add_paragraph("Mode relecture : la situation décrite est celle d'une date passée, à titre d'illustration.")

    for heading, paragraphs in sections(d).items():
        doc.add_heading(heading, level=1)
        for text in paragraphs:
            doc.add_paragraph(text)
        if heading.startswith("1."):
            doc.add_picture(chart_deviation(d), width=Inches(6.0))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if heading.startswith("2."):
            top = d.table.sort_values("Variation_pct", ascending=False).head(8)
            tbl = doc.add_table(rows=1, cols=6)
            tbl.style = "Light Grid Accent 1"
            for i, h in enumerate(["Province", "Prix actuel", "Prix projeté", "Variation", "Écart moy. 5 ans", "Vigilance"]):
                tbl.rows[0].cells[i].text = h
            for r in top.itertuples():
                cells = tbl.add_row().cells
                for i, v in enumerate([r.Province, kg(r.Prix_Actuel), kg(r.Prix_Projete), pct(r.Variation_pct),
                                       pct(r.Ecart_Moy5_pct) if pd.notna(r.Ecart_Moy5_pct) else "n.d.", r.Vigilance]):
                    cells[i].text = str(v)
            doc.add_paragraph()
            trend = chart_trend(d)
            if trend is not None:
                doc.add_picture(trend, width=Inches(6.0))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


__all__ = ["BulletinData", "build_docx", "build_text", "sections", "np"]
